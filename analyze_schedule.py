"""Sequence jobs across machines, compare dispatching rules, and write everything to one Excel file.

Reads the validated CSV produced by validate_data.py, simulates the shop floor
under several dispatching (sequencing) rules, and writes a brand-new workbook
(no template) with these sheets:
  - Input Data        the validated records the analysis was run on
  - Rule Comparison   makespan, tardiness, late jobs, idle time and utilization per rule
  - Best Schedule     job-by-job machine assignment, start/end times and lateness under the best rule
  - Machine Summary   busy hours, idle hours and utilization per machine under the best rule
  - Capacity Plan     what happens to lateness as you add or remove machines
  - Weekly Load       hours of work released each week vs. available capacity
  - Gantt             chart of the best schedule (also saved as a PNG)

Dispatching rules (jobs are non-preemptive; whenever a machine frees up, it picks
the waiting job with the best priority):
  - FIFO   first in, first out: earliest release time
  - SPT    shortest processing time first: clears the queue, minimises average flow time
  - EDD    earliest due date first: classic rule for minimising maximum lateness
  - SLACK  least slack first: due - now - processing time
  - CR     critical ratio: (due - now) / processing time, smallest first
  - WSPT   weighted SPT: highest order value per processing hour first

The "best" rule is the one with the lowest total tardiness (ties broken by makespan).
"""

import argparse
import csv
import os

import matplotlib

matplotlib.use("Agg")  # no display needed; we only save to file
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from scheduling_common import (  # noqa: E402
    CSV_FIELDS,
    HOURS_PER_WEEK,
    OUTPUT_GANTT_PATH,
    OUTPUT_WORKBOOK_PATH,
    VALIDATED_DATA_PATH,
    Job,
)

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LATE_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
BEST_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

# Each rule maps (job, current_time) -> priority; the LOWEST value is dispatched first.
RULES = {
    "FIFO": lambda j, t: (j.release_hour, j.job_id),
    "SPT": lambda j, t: (j.processing_hours, j.job_id),
    "EDD": lambda j, t: (j.due_hour, j.job_id),
    "SLACK": lambda j, t: (j.due_hour - t - j.processing_hours, j.job_id),
    "CR": lambda j, t: ((j.due_hour - t) / j.processing_hours, j.job_id),
    "WSPT": lambda j, t: (-j.order_value / j.processing_hours, j.job_id),
}

SCHEDULE_FIELDS = [
    "job_id",
    "customer",
    "product_family",
    "machine",
    "release_hour",
    "start_hour",
    "end_hour",
    "due_hour",
    "wait_hours",
    "lateness_hours",
    "late",
]


def read_jobs(path: str) -> list[Job]:
    with open(path, newline="", encoding="utf-8") as f:
        return [Job(**row) for row in csv.DictReader(f)]


def simulate(jobs: list[Job], machines: int, rule) -> list[dict]:
    """List-scheduling simulation: whenever the earliest-free machine is idle, dispatch by the rule."""
    free_at = [0.0] * machines
    remaining = list(jobs)
    schedule = []

    while remaining:
        m = min(range(machines), key=lambda i: free_at[i])
        t = free_at[m]
        waiting = [j for j in remaining if j.release_hour <= t]
        if not waiting:  # machine idles until the next order arrives
            t = min(j.release_hour for j in remaining)
            waiting = [j for j in remaining if j.release_hour <= t]

        job = min(waiting, key=lambda j: rule(j, t))
        start, end = t, t + job.processing_hours
        free_at[m] = end
        remaining.remove(job)

        lateness = end - job.due_hour
        schedule.append(
            {
                "job_id": job.job_id,
                "customer": job.customer,
                "product_family": job.product_family,
                "machine": m + 1,
                "release_hour": job.release_hour,
                "start_hour": round(start, 2),
                "end_hour": round(end, 2),
                "due_hour": job.due_hour,
                "wait_hours": round(start - job.release_hour, 2),
                "lateness_hours": round(max(lateness, 0.0), 2),
                "late": lateness > 1e-9,
                "_processing": job.processing_hours,
            }
        )
    return schedule


def schedule_metrics(schedule: list[dict], machines: int) -> dict:
    makespan = max(s["end_hour"] for s in schedule)
    busy = sum(s["_processing"] for s in schedule)
    available = machines * makespan
    late = [s for s in schedule if s["late"]]
    return {
        "makespan_hours": round(makespan, 1),
        "total_tardiness_hours": round(sum(s["lateness_hours"] for s in schedule), 1),
        "max_tardiness_hours": round(max(s["lateness_hours"] for s in schedule), 1),
        "late_jobs": len(late),
        "on_time_pct": round((1 - len(late) / len(schedule)) * 100, 1),
        "avg_flow_time_hours": round(sum(s["end_hour"] - s["release_hour"] for s in schedule) / len(schedule), 1),
        "idle_hours": round(available - busy, 1),
        "utilization_pct": round(busy / available * 100, 1) if available else 0.0,
    }


def compare_rules(jobs: list[Job], machines: int) -> tuple[list[dict], dict[str, list[dict]]]:
    schedules = {name: simulate(jobs, machines, rule) for name, rule in RULES.items()}
    rows = []
    for name, sched in schedules.items():
        rows.append({"rule": name, **schedule_metrics(sched, machines)})
    rows.sort(key=lambda r: (r["total_tardiness_hours"], r["makespan_hours"]))
    for i, row in enumerate(rows):
        row["rank"] = i + 1
    return rows, schedules


def build_capacity_plan(jobs: list[Job], rule_name: str, current: int, max_machines: int) -> list[dict]:
    rows = []
    for m in range(1, max_machines + 1):
        sched = simulate(jobs, m, RULES[rule_name])
        metrics = schedule_metrics(sched, m)
        rows.append(
            {
                "machines": m,
                "current": "<- current" if m == current else "",
                "late_jobs": metrics["late_jobs"],
                "on_time_pct": metrics["on_time_pct"],
                "total_tardiness_hours": metrics["total_tardiness_hours"],
                "makespan_hours": metrics["makespan_hours"],
                "utilization_pct": metrics["utilization_pct"],
            }
        )
    return rows


def build_machine_summary(schedule: list[dict], machines: int) -> list[dict]:
    makespan = max(s["end_hour"] for s in schedule)
    rows = []
    for m in range(1, machines + 1):
        jobs_m = [s for s in schedule if s["machine"] == m]
        busy = sum(s["_processing"] for s in jobs_m)
        rows.append(
            {
                "machine": m,
                "jobs": len(jobs_m),
                "busy_hours": round(busy, 1),
                "idle_hours": round(makespan - busy, 1),
                "utilization_pct": round(busy / makespan * 100, 1),
                "late_jobs": sum(1 for s in jobs_m if s["late"]),
            }
        )
    return rows


def build_weekly_load(jobs: list[Job], machines: int) -> list[dict]:
    last_week = int(max(j.release_hour for j in jobs) // HOURS_PER_WEEK)
    capacity = machines * HOURS_PER_WEEK
    rows = []
    for week in range(last_week + 1):
        week_jobs = [j for j in jobs if week * HOURS_PER_WEEK <= j.release_hour < (week + 1) * HOURS_PER_WEEK]
        hours = sum(j.processing_hours for j in week_jobs)
        rows.append(
            {
                "week": week + 1,
                "jobs_released": len(week_jobs),
                "work_hours_released": round(hours, 1),
                "capacity_hours": capacity,
                "load_pct": round(hours / capacity * 100, 1),
            }
        )
    return rows


def draw_gantt(schedule: list[dict], machines: int, rule_name: str, path: str) -> None:
    fig, ax = plt.subplots(figsize=(12, 1.2 + machines * 0.9))
    for s in schedule:
        color = "#d9534f" if s["late"] else "#4a90c2"
        ax.barh(s["machine"], s["end_hour"] - s["start_hour"], left=s["start_hour"], color=color, edgecolor="white", height=0.6)
        if s["end_hour"] - s["start_hour"] >= 4:
            ax.text((s["start_hour"] + s["end_hour"]) / 2, s["machine"], s["job_id"][-3:], ha="center", va="center", color="white", fontsize=7)
    for week in range(1, int(max(s["end_hour"] for s in schedule) // HOURS_PER_WEEK) + 1):
        ax.axvline(week * HOURS_PER_WEEK, color="grey", linestyle=":", linewidth=0.8)
    ax.set_yticks(range(1, machines + 1))
    ax.set_yticklabels([f"Machine {m}" for m in range(1, machines + 1)])
    ax.invert_yaxis()
    ax.set_xlabel("Working hours (dotted lines = week boundaries)")
    ax.set_title(f"Schedule under {rule_name}  (red = finishes after due date)")
    fig.tight_layout()
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    fig.savefig(path, dpi=130)
    plt.close(fig)


def autosize_columns(ws) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 40)


def write_sheet(ws, fieldnames: list[str], rows: list[dict]) -> None:
    ws.append(fieldnames)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(f, "") for f in fieldnames])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(fieldnames))}{len(rows) + 1}"
    autosize_columns(ws)


def highlight_rows(ws, fieldnames: list[str], column: str, predicate, fill) -> None:
    col = fieldnames.index(column)
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if predicate(row_cells[col].value):
            for cell in row_cells:
                cell.fill = fill


def build_workbook(
    input_rows: list[dict],
    comparison: list[dict],
    best_schedule: list[dict],
    machine_summary: list[dict],
    capacity_plan: list[dict],
    weekly_load: list[dict],
    gantt_path: str,
) -> Workbook:
    wb = Workbook()

    ws_input = wb.active
    ws_input.title = "Input Data"
    write_sheet(ws_input, CSV_FIELDS, input_rows)

    ws_cmp = wb.create_sheet("Rule Comparison")
    cmp_fields = [
        "rank",
        "rule",
        "makespan_hours",
        "total_tardiness_hours",
        "max_tardiness_hours",
        "late_jobs",
        "on_time_pct",
        "avg_flow_time_hours",
        "idle_hours",
        "utilization_pct",
    ]
    write_sheet(ws_cmp, cmp_fields, comparison)
    highlight_rows(ws_cmp, cmp_fields, "rank", lambda v: v == 1, BEST_FILL)

    ws_best = wb.create_sheet("Best Schedule")
    best_sorted = sorted(best_schedule, key=lambda s: (s["machine"], s["start_hour"]))
    write_sheet(ws_best, SCHEDULE_FIELDS, best_sorted)
    highlight_rows(ws_best, SCHEDULE_FIELDS, "late", lambda v: v is True, LATE_FILL)

    ws_mach = wb.create_sheet("Machine Summary")
    write_sheet(ws_mach, ["machine", "jobs", "busy_hours", "idle_hours", "utilization_pct", "late_jobs"], machine_summary)

    ws_cap = wb.create_sheet("Capacity Plan")
    cap_fields = ["machines", "current", "late_jobs", "on_time_pct", "total_tardiness_hours", "makespan_hours", "utilization_pct"]
    write_sheet(ws_cap, cap_fields, capacity_plan)
    highlight_rows(ws_cap, cap_fields, "current", lambda v: v == "<- current", BEST_FILL)

    ws_week = wb.create_sheet("Weekly Load")
    write_sheet(ws_week, ["week", "jobs_released", "work_hours_released", "capacity_hours", "load_pct"], weekly_load)
    highlight_rows(ws_week, ["week", "jobs_released", "work_hours_released", "capacity_hours", "load_pct"], "load_pct", lambda v: isinstance(v, (int, float)) and v > 100, LATE_FILL)

    ws_gantt = wb.create_sheet("Gantt")
    ws_gantt.add_image(XLImage(gantt_path), "A1")

    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare dispatching rules on validated job data and write results to Excel.")
    parser.add_argument("--input", type=str, default=VALIDATED_DATA_PATH, help="Path to the validated CSV")
    parser.add_argument("--output", type=str, default=OUTPUT_WORKBOOK_PATH, help="Path to write the Excel workbook to")
    parser.add_argument("--gantt", type=str, default=OUTPUT_GANTT_PATH, help="Path to write the Gantt chart PNG to")
    parser.add_argument("--machines", type=int, default=3, help="Number of machines / crews available (default: 3)")
    parser.add_argument("--max-machines", type=int, default=None, help="Largest machine count in the capacity plan (default: machines + 3)")
    args = parser.parse_args()

    jobs = read_jobs(args.input)
    with open(args.input, newline="", encoding="utf-8") as f:
        input_rows = list(csv.DictReader(f))

    comparison, schedules = compare_rules(jobs, args.machines)
    best_rule = comparison[0]["rule"]
    best_schedule = schedules[best_rule]

    machine_summary = build_machine_summary(best_schedule, args.machines)
    capacity_plan = build_capacity_plan(jobs, best_rule, args.machines, args.max_machines or args.machines + 3)
    weekly_load = build_weekly_load(jobs, args.machines)

    draw_gantt(best_schedule, args.machines, best_rule, args.gantt)
    wb = build_workbook(input_rows, comparison, best_schedule, machine_summary, capacity_plan, weekly_load, args.gantt)
    os.makedirs(os.path.dirname(args.output) or ".", exist_ok=True)
    wb.save(args.output)

    fifo = next(r for r in comparison if r["rule"] == "FIFO")
    best = comparison[0]
    print(f"Scheduled {len(jobs)} jobs on {args.machines} machine(s) -> {args.output}")
    print(f"Best rule: {best_rule} ({best['late_jobs']} late jobs, {best['total_tardiness_hours']}h total tardiness)")
    print(f"vs. FIFO:  {fifo['late_jobs']} late jobs, {fifo['total_tardiness_hours']}h total tardiness")


if __name__ == "__main__":
    main()
